"""Step 4 · Build BESS_IEX_Financial_Model.xlsx (needs openpyxl, e.g. python3.12 source/mkxlsx.py)."""
import json, os
HERE = os.path.dirname(os.path.abspath(__file__)); ROOT = os.path.dirname(HERE)
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment

EX = json.load(open(os.path.join(HERE, 'export.json')))
PR = json.load(open(os.path.join(HERE, 'prices.json')))
dates = PR['dates'][-365:]; P = {m: PR[m][-365:] for m in ['DAM','GDAM','RTM']}
N = len(dates)
OUT = os.path.join(ROOT, 'BESS_IEX_Financial_Model.xlsx')

F = 'Arial'
fB = Font(name=F, size=10); fBold = Font(name=F, size=10, bold=True)
fIn = Font(name=F, size=10, color='0000FF'); fLink = Font(name=F, size=10, color='008000')
fH1 = Font(name=F, size=14, bold=True); fH2 = Font(name=F, size=11, bold=True, color='FFFFFF'); fNote = Font(name=F, size=9, italic=True, color='555555')
fillKey = PatternFill('solid', fgColor='FFFF00'); fillHead = PatternFill('solid', fgColor='1F3A5F'); fillSub = PatternFill('solid', fgColor='E8EEF4')
thin = Side(style='thin', color='BFBFBF'); bTop = Border(top=Side(style='thin', color='000000'))
CR = '#,##0.0;(#,##0.0);"-"'; CR0 = '#,##0;(#,##0);"-"'; PCT = '0.0%;(0.0%);"-"'; PCT2 = '0.00%;(0.00%);"-"'; NUM = '#,##0;(#,##0);"-"'; NUM2 = '#,##0.00;(#,##0.00);"-"'; X = '0.00"x"'

wb = Workbook()
def sheet(title, first=False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = title; ws.sheet_view.showGridLines = False
    return ws
def style_all(ws):
    for row in ws.iter_rows():
        for c in row:
            if c.font is None or c.font.name != F: c.font = Font(name=F, size=c.font.size or 10, bold=c.font.bold, italic=c.font.italic, color=c.font.color)

# ------------------------------------------------------------------ Cover (filled later)
cov = sheet('Cover', first=True)

# ------------------------------------------------------------------ Assumptions
A = sheet('Assumptions'); REF = {}
A.column_dimensions['A'].width = 46; A.column_dimensions['B'].width = 14; A.column_dimensions['C'].width = 18; A.column_dimensions['D'].width = 90
A['A1'] = 'Assumptions'; A['A1'].font = fH1
A['A2'] = 'Blue = input you can change · yellow = key assumption · black = formula · green = link from another sheet. Percentages are stored as fractions (90% = 0.9).'; A['A2'].font = fNote
r = 4
def head(ws, row, text, ncol=4):
    ws.cell(row=row, column=1, value=text).font = fH2
    for c in range(1, ncol+1): ws.cell(row=row, column=c).fill = fillHead
def inp(name, label, val, unit, fmt=None, note='', key=False, options=None):
    global r
    A.cell(row=r, column=1, value=label).font = fB
    c = A.cell(row=r, column=2, value=val); c.font = fIn
    if fmt: c.number_format = fmt
    if key: c.fill = fillKey
    A.cell(row=r, column=3, value=unit).font = fB
    A.cell(row=r, column=4, value=note).font = fNote
    if options:
        dv = DataValidation(type='list', formula1='"' + ','.join(options) + '"', allow_blank=False); A.add_data_validation(dv); dv.add(c)
    REF[name] = f"Assumptions!$B${r}"; r += 1
def sec(text):
    global r
    r += 1; head(A, r, text); r += 1

sec('Scenario selectors (Solar + BESS model sheet)')
inp('plant', 'Name of solar plant', 'P7', 'code', note='P7 Tumkur · P8 Bidar · P12 Kunnoor 4. Also used by the CATL sheets.', key=True, options=['P7','P8','P12'])
inp('dur', 'BESS duration', 2, 'hours', note='2 or 4 hours of discharge at the full night connectivity.', key=True, options=['2','4'])
inp('case', 'Charging case', 'IEX + T-GNA', 'text', note='Own solar = Solar + BESS project. IEX + T-GNA / IEX + GNA = standalone BESS buying on IEX.', key=True, options=['Own solar','IEX + T-GNA','IEX + GNA'])
inp('solarSell', 'Own-solar case: sell market', 'GDAM', 'market', note='Green energy from own solar is sold on GDAM (HTML default). The own-solar dispatch table was computed with GDAM.', options=['DAM','GDAM','RTM'])
inp('gridBuy', 'IEX cases: buy market', 'RTM', 'market', options=['DAM','GDAM','RTM'])
inp('gridSell', 'IEX cases: sell market', 'GDAM', 'market', options=['DAM','GDAM','RTM'])
inp('c2', 'Charge window start, 2-hour (IEX cases)', 44, 'block 0–95', note='15-minute block index: 44 = 11:00. Charge window length is set by the connectivity limit (Model sheet).')
inp('d2', 'Discharge window start, 2-hour', 76, 'block 0–95', note='76 = 19:00. The own-solar dispatch table uses this default.')
inp('c4', 'Charge window start, 4-hour (IEX cases)', 38, 'block 0–95', note='38 = 09:30.')
inp('d4', 'Discharge window start, 4-hour', 72, 'block 0–95', note='72 = 18:00. The own-solar dispatch table uses this default.')
inp('solarSize', 'Solar sizing, share of average-day charging need', 1.0, '%', PCT, note='Own-solar dispatch table was computed at 100%. Changing it resizes solar capex here but not the dispatch energy; re-export from the HTML model to change both.')
inp('boostOn', 'Booster days on (1 = on, 0 = off)', 0, 'switch', note='Evening discharge sold at the booster price on the booster days.', key=True, options=['0','1'])
inp('boostStart', 'Booster days from (September day)', 1, 'day', note='IEX cases are fully live. The own-solar booster uplift uses the 1–20 Sep dispatch.')
inp('boostDays', 'Number of booster days', 20, 'days')
inp('boostPrice', 'Booster sell price', 15, '₹/kWh', NUM2, note='User input: ₹15/kWh sell price.')
inp('eqOn', 'Equity-funded cost on (1 = on, 0 = off)', 1, 'switch', key=True, options=['0','1'])
inp('eqCost', 'Equity-funded cost', 20, '₹ Cr / yr', CR, note='User input: ₹20 Cr a year, expensed in the P&L (tax-deductible), excluded from DSCR.')
inp('eqYears', '…for the first N years', 2, 'years')

sec('Capex & land')
inp('solarCapex', 'Solar capex', 2.05, '₹ Cr / MWp DC', NUM2, key=True, note='User input.')
inp('bessCapex', 'BESS capex, all-inclusive on nameplate', 0.98, '₹ Cr / MWh', NUM2, key=True, note='User input.')
inp('landRate', 'Land rate', 35000, '₹ / acre', NUM, note='User input (plant table).')
inp('landMode', 'Land rate basis', 'Lease', 'text', note='Lease = annual lease escalating; Buy = one-time purchase in year 0.', options=['Lease','Buy'])
inp('landEsc', 'Lease escalation', 0.05, '% / yr', PCT)
inp('cuf', 'Solar CUF (AC)', 0.30, '%', PCT, note='User input (plant table).')
inp('auxMW', 'Solar auxiliary consumption', 1, 'MW per 300 MWac', NUM2, note='User input: 1 MW for the overall plant, scaled to the sized solar.')

sec('BESS technical')
inp('dod', 'Depth of discharge', 0.90, '%', PCT)
inp('rte', 'Round-trip efficiency (AC-AC)', 0.90, '%', PCT, note='User input.')
inp('deg', 'Capacity fade', 0.015, '% / yr', PCT2, note='User input: 1.5%/yr, linear.')
inp('bessAux', 'HVAC auxiliary, share of discharge', 0.01, '%', PCT, note='User input.')
inp('lineLoss', 'Transformer + line loss to CTU meter', 0.01, '%', PCT)
inp('istsLoss', 'ISTS loss on drawal', 0.03, '%', PCT, note='Assumption.')
inp('avail', 'Availability', 0.99, '%', PCT, note='User input.')
inp('augYear', 'Augmentation year (0 = none)', 12, 'year')
inp('augPct', 'Augmentation cost, share of BESS capex', 0.20, '%', PCT, note='Assumption.')

sec('Transmission & exchange charges')
inp('stoa', 'T-GNA (STOA) charge on energy drawn', 586, '₹ / MWh', NUM, key=True, note='Source: SRPC RTA, billing period May 2026 (billing month July 2026): Karnataka drawee DICs 5,952 MW GNA, ₹235.8 Cr → ×1.10 ÷ (31×96×GNA) = ₹146.4/MW/block (Reg. 11, CERC Sharing Regulations 2020 as amended).')
inp('lta', 'GNA (LTA) charge', 3.96, '₹ lakh / MW / month', NUM2, note='Source: same SRPC statement, ₹235.8 Cr ÷ 5,952 MW. Fixed on BESS MW. CTU charges on selling are zero (user input).')
inp('fee', 'IEX transaction fee, each side', 15, '₹ / MWh', NUM, note='User input.')

sec('Finance & O&M')
inp('life', 'Project life', 25, 'years', note='Timeline has 25 columns; keep ≤ 25.')
inp('tax', 'Corporate tax', 0.2517, '%', PCT2)
inp('esc', 'IEX price escalation', 0.0, '% / yr', PCT)
inp('omSolar', 'Solar O&M', 5, '₹ lakh / MWac / yr', NUM2)
inp('omBess', 'BESS O&M, share of capex', 0.005, '% / yr', PCT2, note='User input.')
inp('omEsc', 'O&M escalation', 0.03, '% / yr', PCT)
inp('solDeg', 'Solar degradation', 0.005, '% / yr', PCT2, note='Included in the own-solar dispatch table.')
inp('debtPct', 'Debt share of capex', 0.70, '%', PCT)
inp('kd', 'Interest rate (Kd)', 0.095, '%', PCT2)
inp('tenor', 'Loan tenor (equal principal)', 15, 'years')
inp('ke', 'Cost of equity (Ke)', 0.14, '%', PCT)
inp('N', 'Days of IEX price data', N, 'days', note=f'IEX MCP 15-min, {dates[0]} to {dates[-1]} (files DAM/GDAM/RTM_15min_2025-09-01_to_2026-09-23.xlsx).')

cat = EX['catl']; w4 = cat['P7|4|stoa']; w2c = cat['P7|2|stoa']
sec('CATL 12,000-cycle BESS · IEX only (CATL sheets)')
inp('cDur', 'BESS duration', 4, 'hours', key=True, options=['2','4'])
inp('cTx', 'Transmission for purchase', 'T-GNA', 'text', options=['T-GNA','GNA'])
inp('cBuy', 'Buy market', 'DAM', 'market', options=['DAM','GDAM','RTM'])
inp('cSell', 'Sell market', 'DAM', 'market', options=['DAM','GDAM','RTM'])
inp('cCapex', 'BESS capex', 1.05, '₹ Cr / MWh', NUM2, key=True)
inp('cCycles', 'Rated cycle life', 12000, 'cycles', NUM, note='User input (CATL 12,000-cycle). Not checked against a CATL datasheet.')
inp('cEol', 'End-of-life capacity', 0.70, '%', PCT, note='Assumption: check against the CATL warranty.')
inp('cCal', 'Calendar fade', 0.003, '% / yr', PCT2, note='Assumption.')
inp('cDod', 'Depth of discharge', 0.90, '%', PCT)
inp('cRte', 'Round-trip efficiency', 0.90, '%', PCT)
inp('cAux', 'HVAC auxiliary', 0.01, '%', PCT)
inp('cAvail', 'Availability', 0.99, '%', PCT)
inp('cRep', 'Cell replacement at end of life, share of capex', 0.35, '%', PCT, note='Assumption.')
for hh, w in [(2, w2c), (4, w4)]:
    inp(f'k1c{hh}', f'1 cycle/day, {hh}-hour: charge start', w['w1'][0]['c'], 'block', note='Best windows from the HTML optimiser (DAM/DAM).')
    inp(f'k1d{hh}', f'1 cycle/day, {hh}-hour: discharge start', w['w1'][0]['d'], 'block')
    for i in (0, 1):
        inp(f'k2c{i+1}{hh}', f'2 cycles/day, {hh}-hour: charge {i+1} start', w['w2'][i]['c'], 'block')
        inp(f'k2d{i+1}{hh}', f'2 cycles/day, {hh}-hour: discharge {i+1} start', w['w2'][i]['d'], 'block')

sec('Plant table')
A.cell(row=r, column=1, value='Plant · location').font = fBold
for j, t in enumerate(['Night connectivity MW', 'Ref. solar MWac', 'Ref. solar MWp DC · acres'], start=2): A.cell(row=r, column=j, value=t).font = fBold
r += 1
PL = [('P7','Tumkur',300,300,419,1200), ('P8','Bidar',300,300,445,1200), ('P12','Kunnoor 4',250,250,371,1000)]
pt0 = r
for code, loc, conn, ac, dc, acres in PL:
    A.cell(row=r, column=1, value=code).font = fBold
    c = A.cell(row=r, column=2, value=conn); c.font = fIn; c.fill = fillKey; c.number_format = NUM
    A.cell(row=r, column=3, value=ac).font = fIn
    A.cell(row=r, column=4, value=f'{dc} MWp DC · {acres} acres · {loc}. Source: user plant table (DC/AC ratio and acres per MWp are used to size solar).').font = fNote
    # hidden helper columns E/F for DC and acres
    A.cell(row=r, column=5, value=dc).font = fIn; A.cell(row=r, column=6, value=acres).font = fIn
    r += 1
A.cell(row=pt0-1, column=5, value='MWp DC').font = fBold; A.cell(row=pt0-1, column=6, value='Acres').font = fBold
PTC = f"Assumptions!$A${pt0}:$A${pt0+2}"
def ptv(col): return f"INDEX(Assumptions!${col}${pt0}:${col}${pt0+2},MATCH({REF['plant']},{PTC},0))"
A.column_dimensions['E'].width = 10; A.column_dimensions['F'].width = 10

# ------------------------------------------------------------------ Prices (96-block average profile, duplicated for window wrap)
PS = sheet('Prices')
PS['A1'] = f'IEX MCP average by 15-minute block, ₹/MWh · {dates[0]} to {dates[-1]} ({N} days)'; PS['A1'].font = fH1
PS['A2'] = 'Row 3 is the block position (0–191; blocks 96–191 repeat 0–95 so windows can run past midnight). Source: IEX DAM / GDAM / RTM 15-minute files provided.'; PS['A2'].font = fNote
PS.cell(row=3, column=1, value='Market / block →').font = fBold
for p in range(192):
    c = PS.cell(row=3, column=2+p, value=p); c.font = fBold
    b = p % 96; PS.cell(row=4, column=2+p, value=f"{b//4:02d}:{(b%4)*15:02d}").font = fNote
for i, m in enumerate(['DAM','GDAM','RTM']):
    avg = [sum(P[m][d][b] for d in range(N))/N for b in range(96)]
    PS.cell(row=5+i, column=1, value=m).font = fBold
    for p in range(192):
        c = PS.cell(row=5+i, column=2+p, value=round(avg[p % 96], 2)); c.font = fB; c.number_format = NUM
PS.column_dimensions['A'].width = 18
PRNG = "Prices!$B$5:$GK$7"; PIDX = "Prices!$B$3:$GK$3"; PMKT = "Prices!$A$5:$A$7"
def win(mkt, s, n): return f"SUMPRODUCT(({PMKT}={mkt})*({PIDX}>={s})*({PIDX}<({s})+({n}))*{PRNG})/({n})"

# ------------------------------------------------------------------ SepDays (daily profiles 1–30 Sep, for booster days)
SD = sheet('SepDays')
SD['A1'] = 'September daily IEX prices, ₹/MWh, for booster-day calculations'; SD['A1'].font = fH1
SD['A2'] = 'One row per market and September day (latest year in the data: 1–23 Sep 2026, 24–30 Sep 2025). Columns repeat blocks 0–95 twice.'; SD['A2'].font = fNote
SD.cell(row=3, column=1, value='Market').font = fBold; SD.cell(row=3, column=2, value='Day / block →').font = fBold
for p in range(192): SD.cell(row=3, column=3+p, value=p).font = fBold
pick = {}
for i, ds in enumerate(dates):
    if ds[5:7] == '09': pick[int(ds[8:10])] = i
rr = 4
for m in ['DAM','GDAM','RTM']:
    for day in range(1, 31):
        SD.cell(row=rr, column=1, value=m).font = fB; SD.cell(row=rr, column=2, value=day).font = fB
        if day in pick:
            row = P[m][pick[day]]
            for p in range(192): SD.cell(row=rr, column=3+p, value=row[p % 96]).font = fB
        rr += 1
SDR = f"SepDays!$C$4:$GL${rr-1}"; SDI = "SepDays!$C$3:$GL$3"; SDM = f"SepDays!$A$4:$A${rr-1}"; SDD = f"SepDays!$B$4:$B${rr-1}"
SEPN = rr - 4

# ------------------------------------------------------------------ Dispatch (own-solar, from the HTML 15-min simulation)
DS = sheet('Dispatch')
DS['A1'] = 'Own-solar dispatch · results of the 15-minute simulation in the HTML model'; DS['A1'].font = fH1
DS['A2'] = ('Hardcoded outputs, per plant | duration | booster (0/1), at default windows (2h: 19:00, 4h: 18:00), GDAM sale, 100% solar sizing and default BESS technical inputs. '
            'Each day the battery charges from the cheapest solar blocks; surplus solar is sold in the day. Includes solar degradation and BESS fade by year. Re-export from the HTML if those inputs change.'); DS['A2'].font = fNote
DS.cell(row=4, column=1, value='Key').font = fBold; DS.cell(row=4, column=2, value='Series').font = fBold
for y in range(1, 26): DS.cell(row=4, column=2+y, value=f'Yr {y}').font = fBold
rr = 5
for key, v in EX['dispatch'].items():
    for s, lab, div in [('bessMWh','BESS MWh sold',1), ('bessRev','BESS revenue ₹ Cr',1e7), ('expMWh','Surplus solar MWh sold',1), ('expRev','Surplus solar revenue ₹ Cr',1e7)]:
        DS.cell(row=rr, column=1, value=key).font = fB; DS.cell(row=rr, column=2, value=s).font = fB
        for y in range(25):
            c = DS.cell(row=rr, column=3+y, value=v[s][y]/div); c.font = fIn; c.number_format = CR if div > 1 else NUM
        DS.cell(row=rr, column=28, value=lab + (f" · {v['fullDays']} of {N} days fully charged in yr 1" if s == 'bessMWh' else '')).font = fNote
        rr += 1
DK = f"Dispatch!$A$5:$A${rr-1}"; DSR = f"Dispatch!$B$5:$B${rr-1}"
DS.column_dimensions['A'].width = 12; DS.column_dimensions['B'].width = 10
def disp(series, ycol_idx):  # ycol_idx 0..24
    col = L(3+ycol_idx)
    return f"SUMPRODUCT(({DK}=$C$KEYROW)*({DSR}=\"{series}\")*Dispatch!${col}$5:${col}${rr-1})"

# ------------------------------------------------------------------ Model builder
YC0 = 4  # column D = year 0; E..AC = years 1..25
def ycol(y): return L(YC0 + y)
class M:
    def __init__(self, ws, title):
        self.ws = ws; self.r = 1; self.ref = {}
        ws.column_dimensions['A'].width = 44; ws.column_dimensions['B'].width = 14; ws.column_dimensions['C'].width = 16
        for y in range(0, 26): ws.column_dimensions[ycol(y)].width = 11
        ws['A1'] = title; ws['A1'].font = fH1; self.r = 3
    def sec(self, t):
        self.r += 1; head(self.ws, self.r, t, ncol=YC0+25); self.r += 1
    def s(self, name, label, formula, unit='', fmt=NUM2, link=False, bold=False):
        ws = self.ws; ws.cell(row=self.r, column=1, value=label).font = fBold if bold else fB
        ws.cell(row=self.r, column=2, value=unit).font = fNote
        c = ws.cell(row=self.r, column=3, value='=' + formula); c.number_format = fmt
        c.font = fLink if link else (fBold if bold else fB)
        self.ref[name] = f"$C${self.r}"; self.r += 1
    def t(self, name, label, fn, unit='₹ Cr', fmt=CR, y0=None, bold=False, total=False):
        ws = self.ws; ws.cell(row=self.r, column=1, value=label).font = fBold if bold else fB
        ws.cell(row=self.r, column=2, value=unit).font = fNote
        self.ref[name] = self.r
        if y0 is not None:
            c = ws.cell(row=self.r, column=YC0, value='=' + y0); c.number_format = fmt; c.font = fBold if bold else fB
        for y in range(1, 26):
            c = ws.cell(row=self.r, column=YC0+y, value='=' + fn(y, ycol(y), ycol(y-1))); c.number_format = fmt; c.font = fBold if bold else fB
            if total: c.border = bTop
        if total and y0 is not None: ws.cell(row=self.r, column=YC0).border = bTop
        self.r += 1
    def R(self, name, col): return f"{col}${self.ref[name]}" if isinstance(self.ref[name], int) else self.ref[name]
    def row(self, name): return self.ref[name]
    def years_header(self):
        ws = self.ws
        ws.cell(row=self.r, column=1, value='Year').font = fBold
        for y in range(0, 26):
            c = ws.cell(row=self.r, column=YC0+y, value=y); c.font = fBold; c.number_format = '0'; c.fill = fillSub
        self.ref['y'] = self.r; self.r += 1

a = lambda k: REF[k]

# ================= Solar + BESS model
ws = sheet('Model'); m = M(ws, 'Solar + BESS model · selected scenario (all ₹ in Crore unless stated)')
ws['A2'] = 'Driven by the Assumptions sheet. Change the plant, duration and charging case there; every figure below recalculates.'; ws['A2'].font = fNote
m.sec('Scenario & sizing from night connectivity')
m.s('conn', 'Night connectivity', ptv('B'), 'MW', NUM, link=True)
m.s('refAC', 'Reference solar AC (plant table)', ptv('C'), 'MWac', NUM, link=True)
m.s('refDC', 'Reference solar DC (plant table)', ptv('E'), 'MWp', NUM, link=True)
m.s('refAcres', 'Reference land (plant table)', ptv('F'), 'acres', NUM, link=True)
m.s('h', 'Duration', a('dur'), 'hours', '0', link=True)
m.s('isSolar', 'Own-solar case (1 = yes)', f'IF({a("case")}="Own solar",1,0)', 'flag', '0')
m.s('isLTA', 'GNA / LTA transmission (1 = yes)', f'IF({a("case")}="IEX + GNA",1,0)', 'flag', '0')
m.s('eta', 'One-way efficiency = √RTE', f'SQRT({a("rte")})', '%', PCT)
m.s('soh1', 'Year-1 average capacity = 1 − fade ÷ 2', f'1-{a("deg")}/2', '%', PCT2)
m.s('kD', 'Discharge chain = η × (1 − HVAC) × (1 − line)', f'{m.R("eta",0)}*(1-{a("bessAux")})*(1-{a("lineLoss")})', '%', PCT)
m.s('sell', 'Energy sold per cycle = connectivity × hours', f'{m.R("conn",0)}*{m.R("h",0)}', 'MWh', NUM, bold=True)
m.s('stored', 'Stored energy needed', f'{m.R("sell",0)}/{m.R("kD",0)}', 'MWh', NUM2)
m.s('name', 'BESS nameplate = stored ÷ DoD ÷ yr-1 capacity', f'{m.R("stored",0)}/{a("dod")}/{m.R("soh1",0)}', 'MWh', NUM, bold=True)
m.s('needSolar', 'Charging need, energy into PCS', f'{m.R("stored",0)}/{m.R("eta",0)}', 'MWh / day', NUM2)
m.s('needGrid', 'IEX purchase per cycle (after ISTS + line loss)', f'{m.R("needSolar",0)}/(1-{a("lineLoss")})/(1-{a("istsLoss")})', 'MWh', NUM2)
m.s('nb', 'Discharge window length', f'{m.R("h",0)}*4', 'blocks', '0')
m.s('nc', 'Charge window length at connectivity limit', f'MIN(96-{m.R("nb",0)},ROUNDUP({m.R("needGrid",0)}/({m.R("conn",0)}*0.25)-0.000000001,0))', 'blocks', '0')
m.s('ac', 'Solar AC capacity (own-solar case)', f'{m.R("isSolar",0)}*MIN({m.R("conn",0)},{a("solarSize")}*{m.R("needSolar",0)}/({a("cuf")}*24))', 'MWac', NUM2)
m.s('dc', 'Solar DC capacity = AC × plant DC/AC', f'{m.R("ac",0)}*{m.R("refDC",0)}/{m.R("refAC",0)}', 'MWp', NUM2)
m.s('acres', 'Land = DC × plant acres per MWp', f'{m.R("dc",0)}*{m.R("refAcres",0)}/{m.R("refDC",0)}', 'acres', NUM)
m.s('aux', 'Solar auxiliary load', f'{a("auxMW")}*{m.R("ac",0)}/300', 'MW', NUM2)
m.sec('IEX prices for the selected windows (live from the Prices sheet)')
m.s('dS', 'Discharge window start', f'IF({m.R("h",0)}=2,{a("d2")},{a("d4")})', 'block', '0', link=True)
m.s('cS', 'Charge window start (IEX cases)', f'IF({m.R("h",0)}=2,{a("c2")},{a("c4")})', 'block', '0', link=True)
m.s('sMkt', 'Sell market', f'IF({m.R("isSolar",0)}=1,{a("solarSell")},{a("gridSell")})', 'market', '@', link=True)
m.s('bMkt', 'Buy market', a('gridBuy'), 'market', '@', link=True)
m.s('winS', 'Average sale price in discharge window', win(m.R('sMkt',0), m.R('dS',0), m.R('nb',0)), '₹ / MWh', NUM)
m.s('winB', 'Average purchase price in charge window', win(m.R('bMkt',0), m.R('cS',0), m.R('nc',0)), '₹ / MWh', NUM)
m.s('allS', 'All-day average of sell market (for solar aux)', f"SUMPRODUCT(({PMKT}={m.R('sMkt',0)})*({PIDX}<96)*{PRNG})/96", '₹ / MWh', NUM)
m.s('nBo', 'Booster days found in data', f"SUMPRODUCT(({SDM}={m.R('sMkt',0)})*({SDD}>={a('boostStart')})*({SDD}<{a('boostStart')}+{a('boostDays')}))", 'days', '0')
m.s('boS', 'Sum of booster-day window averages (actual prices)', f"SUMPRODUCT(({SDM}={m.R('sMkt',0)})*({SDD}>={a('boostStart')})*({SDD}<{a('boostStart')}+{a('boostDays')})*({SDI}>={m.R('dS',0)})*({SDI}<{m.R('dS',0)}+{m.R('nb',0)})*{SDR})/{m.R('nb',0)}", '₹ / MWh', NUM)
m.s('sAvgG', 'Realised sale price, IEX cases (booster applied)', f"IF({a('boostOn')}=1,({a('N')}*{m.R('winS',0)}-{m.R('boS',0)}+{m.R('nBo',0)}*{a('boostPrice')}*1000)/{a('N')},{m.R('winS',0)})", '₹ / MWh', NUM, bold=True)
m.s('ovl', 'Window check (IEX cases)', f'IF({m.R("isSolar",0)}=1,"n/a",IF(OR(MOD({m.R("dS",0)}-{m.R("cS",0)},96)<{m.R("nc",0)},MOD({m.R("cS",0)}-{m.R("dS",0)},96)<{m.R("nb",0)}),"Charge and discharge windows overlap","OK"))', '', '@')
m.s('key', 'Own-solar dispatch key (plant|hours|booster)', f'{a("plant")}&"|"&{m.R("h",0)}&"|"&{a("boostOn")}', '', '@')
KEYROW = m.ref['key'][3:]
m.sec('Capex & financing')
m.s('capB', 'BESS capex = nameplate × ₹ Cr/MWh', f'{m.R("name",0)}*{a("bessCapex")}', '₹ Cr', CR)
m.s('capS', 'Solar capex = MWp × ₹ Cr/MWp', f'{m.R("dc",0)}*{a("solarCapex")}', '₹ Cr', CR)
m.s('land', 'Land value (acres × ₹/acre)', f'{m.R("acres",0)}*{a("landRate")}/10000000', '₹ Cr', CR)
m.s('capL', 'Land purchase in capex', f'IF({a("landMode")}="Buy",{m.R("land",0)},0)', '₹ Cr', CR)
m.s('capex', 'Total capex', f'{m.R("capB",0)}+{m.R("capS",0)}+{m.R("capL",0)}', '₹ Cr', CR, bold=True)
m.s('D', 'Debt share (capped 0–95%)', f'MIN(0.95,MAX(0,{a("debtPct")}))', '%', PCT)
m.s('wacc', 'WACC = Ke × E + Kd × (1 − tax) × D', f'(1-{m.R("D",0)})*{a("ke")}+{m.R("D",0)}*{a("kd")}*(1-{a("tax")})', '%', PCT2, bold=True)
m.s('debt', 'Debt', f'{m.R("capex",0)}*{m.R("D",0)}', '₹ Cr', CR)
m.s('equity', 'Equity', f'{m.R("capex",0)}-{m.R("debt",0)}', '₹ Cr', CR)
m.s('prin', 'Annual principal repayment', f'{m.R("debt",0)}/MAX(1,{a("tenor")})', '₹ Cr', CR)
RS = m.r; m.r += 12  # reserve results block
m.sec('Annual projection')
m.years_header()
Y = lambda c: f"{c}${m.row('y')}"
m.t('flag', 'In project life (1/0)', lambda y,c,p: f"IF({Y(c)}<={a('life')},1,0)", 'flag', '0')
m.t('age', 'Years since augmentation (mid-year)', lambda y,c,p: f"IF(AND({a('augYear')}>0,{Y(c)}>{a('augYear')}),{Y(c)}-{a('augYear')}-0.5,{Y(c)}-0.5)", 'years', '0.0')
m.t('soh', 'BESS capacity, % of nameplate', lambda y,c,p: f"MAX(0.4,1-{a('deg')}*{m.R('age',c)})", '%', PCT)
m.t('bf', 'Energy factor vs year 1', lambda y,c,p: f"{m.R('soh',c)}/{m.ref['soh1']}", 'x', '0.000')
m.t('pe', 'Price escalation index', lambda y,c,p: f"(1+{a('esc')})^({Y(c)}-1)", 'x', '0.000')
m.t('oe', 'O&M escalation index', lambda y,c,p: f"(1+{a('omEsc')})^({Y(c)}-1)", 'x', '0.000')
m.t('bMWh', 'BESS energy sold', lambda y,c,p: f"{m.R('flag',c)}*IF({m.ref['isSolar']}=1,{disp('bessMWh',y-1).replace('$C$KEYROW','$C$'+KEYROW)},{m.ref['sell']}*{a('N')}*{a('avail')}*{m.R('bf',c)})", 'MWh', NUM)
m.t('xMWh', 'Surplus solar sold in the day', lambda y,c,p: f"{m.R('flag',c)}*{m.ref['isSolar']}*{disp('expMWh',y-1).replace('$C$KEYROW','$C$'+KEYROW)}", 'MWh', NUM)
m.t('gMWh', 'IEX energy bought', lambda y,c,p: f"{m.R('flag',c)}*(1-{m.ref['isSolar']})*{m.ref['needGrid']}*{a('N')}*{a('avail')}*{m.R('bf',c)}", 'MWh', NUM)
m.t('bRev', 'BESS evening sales', lambda y,c,p: f"{m.R('flag',c)}*{m.R('pe',c)}*IF({m.ref['isSolar']}=1,{disp('bessRev',y-1).replace('$C$KEYROW','$C$'+KEYROW)},{m.ref['sell']}*{a('N')}*{a('avail')}*{m.R('bf',c)}*{m.ref['sAvgG']}/10000000)")
m.t('xRev', 'Surplus solar sales', lambda y,c,p: f"{m.R('flag',c)}*{m.R('pe',c)}*{m.ref['isSolar']}*{disp('expRev',y-1).replace('$C$KEYROW','$C$'+KEYROW)}")
m.t('rev', 'Revenue', lambda y,c,p: f"{m.R('bRev',c)}+{m.R('xRev',c)}", bold=True, total=True)
m.t('buy', 'IEX purchase', lambda y,c,p: f"{m.R('gMWh',c)}*{m.ref['winB']}/10000000*{m.R('pe',c)}")
m.t('ctu', 'CTU charges (T-GNA per MWh drawn, or GNA fixed)', lambda y,c,p: f"{m.R('flag',c)}*(1-{m.ref['isSolar']})*IF({m.ref['isLTA']}=1,{m.ref['conn']}*{a('lta')}*100000*12/10000000,{m.R('gMWh',c)}*{a('stoa')}/10000000)")
m.t('fees', 'IEX transaction fees', lambda y,c,p: f"({m.R('gMWh',c)}+{m.R('bMWh',c)}+{m.R('xMWh',c)})*{a('fee')}/10000000")
m.t('auxc', 'Solar auxiliary consumption', lambda y,c,p: f"{m.R('flag',c)}*{m.ref['isSolar']}*{m.ref['aux']}*24*{a('N')}*{m.ref['allS']}/10000000*{m.R('pe',c)}")
m.t('om', 'O&M', lambda y,c,p: f"{m.R('flag',c)}*({m.ref['capB']}*{a('omBess')}+{m.ref['isSolar']}*{m.ref['ac']}*{a('omSolar')}*100000/10000000)*{m.R('oe',c)}")
m.t('lease', 'Land lease', lambda y,c,p: f"{m.R('flag',c)}*IF(AND({m.ref['isSolar']}=1,{a('landMode')}=\"Lease\"),{m.ref['land']}*(1+{a('landEsc')})^({Y(c)}-1),0)")
m.t('eqc', 'Equity-funded cost', lambda y,c,p: f"{m.R('flag',c)}*IF(AND({a('eqOn')}=1,{Y(c)}<={a('eqYears')}),{a('eqCost')},0)")
m.t('ebitda', 'EBITDA', lambda y,c,p: f"{m.R('rev',c)}-{m.R('buy',c)}-{m.R('ctu',c)}-{m.R('fees',c)}-{m.R('auxc',c)}-{m.R('om',c)}-{m.R('lease',c)}-{m.R('eqc',c)}", bold=True, total=True)
m.t('aug', 'Augmentation capex', lambda y,c,p: f"{m.R('flag',c)}*IF(AND({a('augYear')}>0,{Y(c)}={a('augYear')}),{m.ref['capB']}*{a('augPct')},0)")
m.t('dep', 'Depreciation (straight line)', lambda y,c,p: f"{m.R('flag',c)}*{m.ref['capex']}/{a('life')}")
m.t('taxP', 'Tax, project basis', lambda y,c,p: f"MAX(0,{m.R('ebitda',c)}-{m.R('dep',c)}-{m.R('aug',c)})*{a('tax')}")
m.t('cfPre', 'Project cash flow, pre-tax', lambda y,c,p: f"{m.R('ebitda',c)}-{m.R('aug',c)}", y0=f"-{m.ref['capex']}")
m.t('cf', 'Project cash flow, post-tax', lambda y,c,p: f"{m.R('ebitda',c)}-{m.R('aug',c)}-{m.R('taxP',c)}", y0=f"-{m.ref['capex']}", bold=True, total=True)
m.ref['close'] = m.r + 3
m.t('open', 'Debt opening balance', lambda y,c,p: (m.ref['debt'] if y == 1 else f"{m.R('close',p)}"))
CLOSE_EXPECT = m.ref['close']
m.t('intr', 'Interest', lambda y,c,p: f"{m.R('open',c)}*{a('kd')}")
m.t('rep', 'Principal repayment', lambda y,c,p: f"IF({Y(c)}<={a('tenor')},MIN({m.R('open',c)},{m.ref['prin']}),0)")
m.t('close', 'Debt closing balance', lambda y,c,p: f"{m.R('open',c)}-{m.R('rep',c)}")
m.t('taxE', 'Tax, after interest', lambda y,c,p: f"MAX(0,{m.R('ebitda',c)}-{m.R('dep',c)}-{m.R('aug',c)}-{m.R('intr',c)})*{a('tax')}")
m.t('ecf', 'Equity cash flow', lambda y,c,p: f"{m.R('ebitda',c)}-{m.R('aug',c)}-{m.R('taxE',c)}-{m.R('intr',c)}-{m.R('rep',c)}", y0=f"-{m.ref['equity']}", bold=True, total=True)
m.t('dscr', 'DSCR (excl. augmentation and equity-funded cost)', lambda y,c,p: f"IF({m.R('intr',c)}+{m.R('rep',c)}>0,({m.R('ebitda',c)}+{m.R('eqc',c)}-{m.R('taxE',c)})/({m.R('intr',c)}+{m.R('rep',c)}),\"\")", 'x', X)
m.t('df', 'Discount factor at WACC', lambda y,c,p: f"1/(1+{m.ref['wacc']})^{Y(c)}", 'x', '0.000')
m.t('pvc', 'PV of costs net of surplus solar', lambda y,c,p: f"({m.R('buy',c)}+{m.R('ctu',c)}+{m.R('fees',c)}+{m.R('auxc',c)}+{m.R('om',c)}+{m.R('lease',c)}+{m.R('eqc',c)}+{m.R('aug',c)}-{m.R('xRev',c)})*{m.R('df',c)}")
m.t('pvm', 'PV of BESS MWh sold', lambda y,c,p: f"{m.R('bMWh',c)}*{m.R('df',c)}", 'MWh', NUM)
m.t('pvr', 'PV of BESS revenue', lambda y,c,p: f"{m.R('bRev',c)}*{m.R('df',c)}")
m.t('cum', 'Cumulative post-tax cash flow', lambda y,c,p: f"{m.R('cum',p)}+{m.R('cf',c)}", y0=f"{m.R('cf','D')}")
m.t('pbk', 'Payback helper', lambda y,c,p: f"IF(AND({m.R('cum',p)}<0,{m.R('cum',c)}>=0),{Y(c)}-1+(-{m.R('cum',p)})/{m.R('cf',c)},0)", 'years', '0.00')
def results(m, RS, extra):
    ws = m.ws; rr0 = RS; head(ws, rr0, 'Results', ncol=YC0+25)
    rng = lambda n: f"$D${m.row(n)}:$AC${m.row(n)}"; rng1 = lambda n: f"$E${m.row(n)}:$AC${m.row(n)}"
    items = [('irr', 'Project IRR, post-tax', f"IFERROR(IRR({rng('cf')}),\"< 0%\")", PCT),
             ('irrPre', 'Project IRR, pre-tax', f"IFERROR(IRR({rng('cfPre')}),\"< 0%\")", PCT),
             ('eirr', 'Equity IRR, post-tax', f"IFERROR(IRR({rng('ecf')}),\"< 0%\")", PCT),
             ('npv', 'NPV at WACC, post-tax', f"$D${m.row('cf')}+NPV({m.ref['wacc']},{rng1('cf')})", CR),
             ('pb', 'Payback, post-tax (years)', f"IF(SUM({rng1('pbk')})>0,SUM({rng1('pbk')}),\"beyond life\")", '0.0'),
             ('dmin', 'Minimum DSCR', f"MIN({rng1('dscr')})", X)] + extra
    for i, (nm, lab, f, fmt) in enumerate(items):
        ws.cell(row=rr0+1+i, column=1, value=lab).font = fBold
        c = ws.cell(row=rr0+1+i, column=3, value='=' + f); c.number_format = fmt; c.font = fBold
        m.ref[nm] = f"$C${rr0+1+i}"
results(m, RS, [('lcos', 'Levelised cost per kWh sold (at WACC)', f"({m.ref['capex']}+SUM($E${m.row('pvc')}:$AC${m.row('pvc')}))*10000000/(SUM($E${m.row('pvm')}:$AC${m.row('pvm')})*1000)", '"₹"0.00'),
                ('lsp', 'Levelised sale price per kWh (at WACC)', f"SUM($E${m.row('pvr')}:$AC${m.row('pvr')})*10000000/(SUM($E${m.row('pvm')}:$AC${m.row('pvm')})*1000)", '"₹"0.00'),
                ('tariff', 'Year-1 average BESS sale tariff', f"IFERROR($E${m.row('bRev')}*10000000/($E${m.row('bMWh')}*1000),0)", '"₹"0.00" /kWh"')])
ws.freeze_panes = 'D5'
MAIN = m

# ================= CATL sheets
def catl(title, cpd):
    ws = sheet(title); m = M(ws, f'CATL 12,000-cycle BESS · IEX only · {cpd} cycle{"s" if cpd > 1 else ""} a day (₹ Crore unless stated)')
    ws['A2'] = 'Uses the plant, losses, charges, tax and financing from Assumptions, and the CATL inputs. No solar.'; ws['A2'].font = fNote
    m.sec('Sizing')
    m.s('conn', 'Night connectivity', ptv('B'), 'MW', NUM, link=True)
    m.s('h', 'Duration', a('cDur'), 'hours', '0', link=True)
    m.s('cpd', 'Cycles a day', str(cpd), 'cycles', '0')
    m.s('eta', 'One-way efficiency', f'SQRT({a("cRte")})', '%', PCT)
    m.s('fpc', 'Fade per cycle = (1 − end-of-life) ÷ rated cycles', f'(1-{a("cEol")})/{a("cCycles")}', '%', '0.00000%')
    m.s('cpy', 'Cycles a year', f'{a("N")}*{a("cAvail")}*{m.R("cpd",0)}', 'cycles', NUM)
    m.s('soh1', 'Year-1 average capacity', f'1-{m.R("fpc",0)}*{m.R("cpy",0)}/2-{a("cCal")}/2', '%', PCT2)
    m.s('kD', 'Discharge chain', f'{m.R("eta",0)}*(1-{a("cAux")})*(1-{a("lineLoss")})', '%', PCT)
    m.s('sell', 'Energy sold per cycle', f'{m.R("conn",0)}*{m.R("h",0)}', 'MWh', NUM, bold=True)
    m.s('stored', 'Stored energy needed', f'{m.R("sell",0)}/{m.R("kD",0)}', 'MWh', NUM2)
    m.s('name', 'Nameplate', f'{m.R("stored",0)}/{a("cDod")}/{m.R("soh1",0)}', 'MWh', NUM, bold=True)
    m.s('needGrid', 'IEX purchase per cycle', f'{m.R("stored",0)}/{m.R("eta",0)}/(1-{a("lineLoss")})/(1-{a("istsLoss")})', 'MWh', NUM2)
    m.s('nb', 'Discharge window length', f'{m.R("h",0)}*4', 'blocks', '0')
    m.s('nc', 'Charge window length', f'MIN(48-{m.R("nb",0)},ROUNDUP({m.R("needGrid",0)}/({m.R("conn",0)}*0.25)-0.000000001,0))', 'blocks', '0')
    m.s('rateY', 'Capacity loss a year = fade per cycle × cycles + calendar', f'{m.R("fpc",0)}*{m.R("cpy",0)}+{a("cCal")}', '%', PCT2)
    m.s('yEol', 'Years to end of life', f'(1-{a("cEol")})/{m.R("rateY",0)}', 'years', '0.0')
    m.sec('Windows & prices')
    wins = []
    if cpd == 1: wins = [('k1c', 'k1d')]
    else: wins = [('k2c1', 'k2d1'), ('k2c2', 'k2d2')]
    ss, bb = [], []
    for i, (kc, kd) in enumerate(wins, start=1):
        m.s(f'c{i}', f'Charge {i} start', f'IF({m.R("h",0)}=2,{a(kc+"2")},{a(kc+"4")})', 'block', '0', link=True)
        m.s(f'd{i}', f'Discharge {i} start', f'IF({m.R("h",0)}=2,{a(kd+"2")},{a(kd+"4")})', 'block', '0', link=True)
        m.s(f'ws{i}', f'Average sale price, discharge {i}', win(a('cSell'), m.R(f'd{i}',0), m.R('nb',0)), '₹ / MWh', NUM)
        m.s(f'wb{i}', f'Average purchase price, charge {i}', win(a('cBuy'), m.R(f'c{i}',0), m.R('nc',0)), '₹ / MWh', NUM)
        ss.append(m.R(f'ws{i}',0)); bb.append(m.R(f'wb{i}',0))
    m.s('sumS', 'Sum of cycle sale prices', '+'.join(ss), '₹ / MWh', NUM)
    m.s('sumB', 'Sum of cycle purchase prices', '+'.join(bb), '₹ / MWh', NUM)
    m.s('b1', 'Year-1 energy bought', f'{m.R("needGrid",0)}*{a("N")}*{a("cAvail")}*{m.R("cpd",0)}', 'MWh', NUM)
    m.s('s1', 'Year-1 energy sold', f'{m.R("sell",0)}*{a("N")}*{a("cAvail")}*{m.R("cpd",0)}', 'MWh', NUM)
    m.s('bc1', 'Year-1 purchase cost', f'{m.R("needGrid",0)}*{a("N")}*{a("cAvail")}*{m.R("sumB",0)}/10000000', '₹ Cr', CR)
    m.s('sr1', 'Year-1 sales revenue', f'{m.R("sell",0)}*{a("N")}*{a("cAvail")}*{m.R("sumS",0)}/10000000', '₹ Cr', CR)
    m.sec('Capex & financing')
    m.s('capB', 'BESS capex', f'{m.R("name",0)}*{a("cCapex")}', '₹ Cr', CR)
    m.s('capex', 'Total capex', m.R('capB',0), '₹ Cr', CR, bold=True)
    m.s('D', 'Debt share', f'MIN(0.95,MAX(0,{a("debtPct")}))', '%', PCT)
    m.s('wacc', 'WACC', f'(1-{m.R("D",0)})*{a("ke")}+{m.R("D",0)}*{a("kd")}*(1-{a("tax")})', '%', PCT2, bold=True)
    m.s('debt', 'Debt', f'{m.R("capex",0)}*{m.R("D",0)}', '₹ Cr', CR)
    m.s('equity', 'Equity', f'{m.R("capex",0)}-{m.R("debt",0)}', '₹ Cr', CR)
    m.s('prin', 'Annual principal repayment', f'{m.R("debt",0)}/MAX(1,{a("tenor")})', '₹ Cr', CR)
    RS = m.r; m.r += 9
    m.sec('Annual projection'); m.years_header()
    Y = lambda c: f"{c}${m.row('y')}"
    m.t('flag', 'In project life (1/0)', lambda y,c,p: f"IF({Y(c)}<={a('life')},1,0)", 'flag', '0')
    m.ref['repc'] = m.r + 14
    m.t('a', 'Years since last cell replacement (start of year)', lambda y,c,p: ('0' if y == 1 else f"IF({m.R('repc',p)}>0,0,{m.R('a',p)}+1)"), 'years', '0')
    m.t('soh', 'Capacity, mid-year', lambda y,c,p: f"MAX(0.3,1-{m.ref['rateY']}*({m.R('a',c)}+0.5))", '%', PCT)
    m.t('bf', 'Energy factor vs year 1', lambda y,c,p: f"{m.R('soh',c)}/{m.ref['soh1']}", 'x', '0.000')
    m.t('pe', 'Price escalation index', lambda y,c,p: f"(1+{a('esc')})^({Y(c)}-1)", 'x', '0.000')
    m.t('oe', 'O&M escalation index', lambda y,c,p: f"(1+{a('omEsc')})^({Y(c)}-1)", 'x', '0.000')
    m.t('sMWh', 'Energy sold', lambda y,c,p: f"{m.R('flag',c)}*{m.ref['s1']}*{m.R('bf',c)}", 'MWh', NUM)
    m.t('gMWh', 'Energy bought', lambda y,c,p: f"{m.R('flag',c)}*{m.ref['b1']}*{m.R('bf',c)}", 'MWh', NUM)
    m.t('rev', 'Revenue', lambda y,c,p: f"{m.R('flag',c)}*{m.ref['sr1']}*{m.R('bf',c)}*{m.R('pe',c)}", bold=True, total=True)
    m.t('buy', 'IEX purchase', lambda y,c,p: f"{m.R('flag',c)}*{m.ref['bc1']}*{m.R('bf',c)}*{m.R('pe',c)}")
    m.t('ctu', 'CTU charges', lambda y,c,p: f"{m.R('flag',c)}*IF({a('cTx')}=\"T-GNA\",{m.R('gMWh',c)}*{a('stoa')}/10000000,{m.ref['conn']}*{a('lta')}*100000*12/10000000)")
    m.t('fees', 'IEX transaction fees', lambda y,c,p: f"({m.R('gMWh',c)}+{m.R('sMWh',c)})*{a('fee')}/10000000")
    m.t('om', 'O&M', lambda y,c,p: f"{m.R('flag',c)}*{m.ref['capB']}*{a('omBess')}*{m.R('oe',c)}")
    m.t('eqc', 'Equity-funded cost', lambda y,c,p: f"{m.R('flag',c)}*IF(AND({a('eqOn')}=1,{Y(c)}<={a('eqYears')}),{a('eqCost')},0)")
    m.t('ebitda', 'EBITDA', lambda y,c,p: f"{m.R('rev',c)}-{m.R('buy',c)}-{m.R('ctu',c)}-{m.R('fees',c)}-{m.R('om',c)}-{m.R('eqc',c)}", bold=True, total=True)
    assert m.r == m.ref['repc'], (m.r, m.ref['repc'])
    m.t('repc', 'Cell replacement at end of life', lambda y,c,p: f"{m.R('flag',c)}*IF(AND(1-{m.ref['rateY']}*({m.R('a',c)}+1)<{a('cEol')},{Y(c)}<{a('life')}),{m.ref['capB']}*{a('cRep')},0)")
    m.t('dep', 'Depreciation', lambda y,c,p: f"{m.R('flag',c)}*{m.ref['capex']}/{a('life')}")
    m.t('taxP', 'Tax, project basis', lambda y,c,p: f"MAX(0,{m.R('ebitda',c)}-{m.R('dep',c)}-{m.R('repc',c)})*{a('tax')}")
    m.t('cfPre', 'Project cash flow, pre-tax', lambda y,c,p: f"{m.R('ebitda',c)}-{m.R('repc',c)}", y0=f"-{m.ref['capex']}")
    m.t('cf', 'Project cash flow, post-tax', lambda y,c,p: f"{m.R('ebitda',c)}-{m.R('repc',c)}-{m.R('taxP',c)}", y0=f"-{m.ref['capex']}", bold=True, total=True)
    m.ref['close'] = m.r + 3
    m.t('open', 'Debt opening balance', lambda y,c,p: (m.ref['debt'] if y == 1 else f"{m.R('close',p)}"))
    CLOSE_EXPECT = m.ref['close']
    m.t('intr', 'Interest', lambda y,c,p: f"{m.R('open',c)}*{a('kd')}")
    m.t('rep', 'Principal repayment', lambda y,c,p: f"IF({Y(c)}<={a('tenor')},MIN({m.R('open',c)},{m.ref['prin']}),0)")
    m.t('close', 'Debt closing balance', lambda y,c,p: f"{m.R('open',c)}-{m.R('rep',c)}")
    m.t('taxE', 'Tax, after interest', lambda y,c,p: f"MAX(0,{m.R('ebitda',c)}-{m.R('dep',c)}-{m.R('repc',c)}-{m.R('intr',c)})*{a('tax')}")
    m.t('ecf', 'Equity cash flow', lambda y,c,p: f"{m.R('ebitda',c)}-{m.R('repc',c)}-{m.R('taxE',c)}-{m.R('intr',c)}-{m.R('rep',c)}", y0=f"-{m.ref['equity']}", bold=True, total=True)
    m.t('dscr', 'DSCR (excl. replacement and equity-funded cost)', lambda y,c,p: f"IF({m.R('intr',c)}+{m.R('rep',c)}>0,({m.R('ebitda',c)}+{m.R('eqc',c)}-{m.R('taxE',c)})/({m.R('intr',c)}+{m.R('rep',c)}),\"\")", 'x', X)
    m.t('cum', 'Cumulative post-tax cash flow', lambda y,c,p: f"{m.R('cum',p)}+{m.R('cf',c)}", y0=f"{m.R('cf','D')}")
    m.t('pbk', 'Payback helper', lambda y,c,p: f"IF(AND({m.R('cum',p)}<0,{m.R('cum',c)}>=0),{Y(c)}-1+(-{m.R('cum',p)})/{m.R('cf',c)},0)", 'years', '0.00')
    results(m, RS, [('repy', 'First cell replacement year', f"IFERROR(INDEX($E${m.row('y')}:$AC${m.row('y')},MATCH(TRUE,INDEX($E${m.row('repc')}:$AC${m.row('repc')}>0,0),0)),\"none\")", '0')])
    ws.freeze_panes = 'D5'
    return m
C1 = catl('CATL 1 cycle', 1); C2 = catl('CATL 2 cycles', 2)

# ------------------------------------------------------------------ Scenarios (reference values from the HTML model)
SC = sheet('Scenarios')
SC['A1'] = 'Reference results from the HTML model (default inputs)'; SC['A1'].font = fH1
SC['A2'] = 'Static values for checking the live sheets. Defaults: solar ₹2.05 Cr/MWp, BESS ₹0.98 Cr/MWh, IEX cases buy RTM / sell GDAM, T-GNA ₹586/MWh, equity-funded cost on unless shown.'; SC['A2'].font = fNote
hdr = ['Plant','Hours','Case','Booster','Equity cost','Project IRR','Equity IRR','NPV ₹ Cr','Capex ₹ Cr','Min DSCR','Nameplate MWh','Solar MWac','BESS MWh sold, yr 1','EBITDA yr 1 ₹ Cr','Levelised cost ₹/kWh']
for j, h in enumerate(hdr, start=1): c = SC.cell(row=4, column=j, value=h); c.font = fH2; c.fill = fillHead; SC.column_dimensions[L(j)].width = 14
rr = 5
for x in EX['results']:
    vals = [x['pk'], x['h'], x['lab'], 'On' if x['bo'] else 'Off', 'On' if x['eq'] else 'Off', x['irr'] if x['irr'] is not None else '< 0%', x['eirr'] if x['eirr'] is not None else '< 0%', x['npv'], x['capex'], x['dscr'], x['name'], x['ac'] if x['lab']=='Own solar' else 0, x['sold'], x['ebitda1'], x['lcos']]
    fm = ['@','0','@','@','@',PCT,PCT,CR,CR,X,NUM,NUM2,NUM,CR,'0.00']
    for j, v in enumerate(vals, start=1): c = SC.cell(row=rr, column=j, value=v); c.font = fB; c.number_format = fm[j-1]
    rr += 1
rr += 1; SC.cell(row=rr, column=1, value='CATL 12,000-cycle · IEX only (DAM/DAM, equity cost on)').font = fBold; rr += 1
for j, h in enumerate(['Plant','Hours','Transmission','1-cycle project IRR','1-cycle equity IRR','2-cycle project IRR','2-cycle equity IRR','2-cycle replacement year'], start=1):
    c = SC.cell(row=rr, column=j, value=h); c.font = fH2; c.fill = fillHead
rr += 1
for k, v in EX['catl'].items():
    pk, h, tx = k.split('|'); vals = [pk, int(h), 'T-GNA' if tx == 'stoa' else 'GNA', v['r1']['irr'], v['r1']['eirr'], v['r2']['irr'], v['r2']['eirr'], ', '.join(map(str, v['r2']['reps'])) or 'none']
    for j, val in enumerate(vals, start=1):
        c = SC.cell(row=rr, column=j, value=val if val is not None else '< 0%'); c.font = fB
        if j in (4,5,6,7): c.number_format = PCT
    rr += 1

# ------------------------------------------------------------------ Cover
cov.column_dimensions['A'].width = 42; cov.column_dimensions['B'].width = 18; cov.column_dimensions['C'].width = 18; cov.column_dimensions['D'].width = 18; cov.column_dimensions['E'].width = 60
cov['A1'] = 'BESS IEX Arbitrage · Financial Model'; cov['A1'].font = Font(name=F, size=16, bold=True)
cov['A2'] = f'P7 Tumkur · P8 Bidar · P12 Kunnoor 4 · BESS sized from CTU night connectivity · IEX prices {dates[0]} to {dates[-1]}'; cov['A2'].font = fNote
rows = [('Sheet', 'What it holds'), ('Assumptions', 'Every input with units and sources. Scenario selectors at the top.'),
        ('Model', 'Solar + BESS model for the selected plant, duration and charging case: sizing, prices, capex, 25-year P&L, cash flow, debt, IRR, NPV, DSCR, levelised cost.'),
        ('CATL 1 cycle / CATL 2 cycles', 'CATL 12,000-cycle BESS on IEX only, one or two cycles a day, with cycle-based fade and cell replacement.'),
        ('Prices / SepDays', 'IEX 15-minute average price profiles and September daily prices (for booster days).'),
        ('Dispatch', 'Own-solar daily charging results from the HTML model’s 15-minute simulation.'),
        ('Scenarios', 'Reference results from the HTML model for checking.')]
for i, (x, y) in enumerate(rows):
    cov.cell(row=4+i, column=1, value=x).font = fBold if i else fH2; cov.cell(row=4+i, column=2, value=y).font = fB if i else fH2
    if i == 0:
        for cc in range(1, 6): cov.cell(row=4, column=cc).fill = fillHead
cov['A12'] = 'Colour legend'; cov['A12'].font = fBold
leg = [('Blue text', 'Input you can change', fIn, None), ('Yellow fill', 'Key assumption', fIn, fillKey), ('Black text', 'Formula', fB, None), ('Green text', 'Link from another sheet', fLink, None)]
for i, (x, y, fo, fi) in enumerate(leg):
    c = cov.cell(row=13+i, column=1, value=x); c.font = fo
    if fi: c.fill = fi
    cov.cell(row=13+i, column=2, value=y).font = fB
cov['A18'] = 'Live results'; cov['A18'].font = fH2
for cc in range(1, 6): cov.cell(row=18, column=cc).fill = fillHead
cov['B18'] = 'Solar + BESS'; cov['C18'] = 'CATL 1 cycle'; cov['D18'] = 'CATL 2 cycles'
for x in ('B18','C18','D18'): cov[x].font = fH2
res = [('Scenario', f"=Assumptions!$B${REF['plant'].split('$')[-1]}&\" · \"&Assumptions!$B${REF['dur'].split('$')[-1]}&\"h · \"&Assumptions!$B${REF['case'].split('$')[-1]}",
        f"=Assumptions!$B${REF['plant'].split('$')[-1]}&\" · \"&Assumptions!$B${REF['cDur'].split('$')[-1]}&\"h · \"&Assumptions!$B${REF['cTx'].split('$')[-1]}", None, '@'),
       ('Project IRR, post-tax', 'irr', 'irr', 'irr', PCT), ('Project IRR, pre-tax', 'irrPre', 'irrPre', 'irrPre', PCT), ('Equity IRR, post-tax', 'eirr', 'eirr', 'eirr', PCT),
       ('WACC', 'wacc', 'wacc', 'wacc', PCT2), ('NPV at WACC, ₹ Cr', 'npv', 'npv', 'npv', CR), ('Payback, years', 'pb', 'pb', 'pb', '0.0'), ('Minimum DSCR', 'dmin', 'dmin', 'dmin', X),
       ('Total capex, ₹ Cr', 'capex', 'capex', 'capex', CR), ('BESS nameplate, MWh', 'name', 'name', 'name', NUM)]
for i, (lab, x1, x2, x3, fm) in enumerate(res):
    rrow = 19+i; cov.cell(row=rrow, column=1, value=lab).font = fBold
    for j, (mm, key, sh) in enumerate([(MAIN, x1, 'Model'), (C1, x2, "'CATL 1 cycle'"), (C2, x3, "'CATL 2 cycles'")], start=2):
        if key is None: continue
        v = key if isinstance(key, str) and key.startswith('=') else f"={sh}!{mm.ref[key]}"
        c = cov.cell(row=rrow, column=j, value=v); c.font = fLink; c.number_format = fm
cov['A31'] = ('How the model works: the BESS is sized so the full night connectivity is exported for the whole window (2 or 4 h). Own-solar charging builds only enough solar to charge that BESS. '
              'IEX-purchase and CATL cases are fully live formulas. Own-solar daily charging (how full the battery gets each day, surplus solar) comes from the HTML model’s 15-minute simulation (Dispatch sheet) at default windows and sizing.')
cov['A31'].font = fNote; cov['A31'].alignment = Alignment(wrap_text=True, vertical='top'); cov.merge_cells('A31:E34'); cov.row_dimensions[31].height = 30
cov['A36'] = 'Units: ₹ Cr = ₹ crore (10 million). Prices ₹/MWh unless shown as ₹/kWh. Percentages are stored as fractions.'; cov['A36'].font = fNote

wb.calculation.fullCalcOnLoad = True
wb.save(OUT); print('saved', OUT)
