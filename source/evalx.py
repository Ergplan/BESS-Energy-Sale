"""Optional · Recalculate the Excel with the `formulas` engine and report errors (pip install formulas). python3.12 source/evalx.py"""
import formulas, sys, time, json, re
from openpyxl import load_workbook
import os
HERE=os.path.dirname(os.path.abspath(__file__)); SRC=os.path.join(os.path.dirname(HERE),'BESS_IEX_Financial_Model.xlsx')
def run(overrides, tag):
    wb=load_workbook(SRC); A=wb['Assumptions']
    lab={}
    for r in range(1,A.max_row+1): lab.setdefault(A.cell(row=r,column=1).value, r)
    for k,v in overrides.items(): A.cell(row=lab[k],column=2).value=v
    tmp=os.path.join(HERE,f'.x_{tag}.xlsx'); wb.save(tmp)
    t0=time.time(); xl=formulas.ExcelModel().loads(tmp).finish(); sol=xl.calculate()
    base=f"'[.x_{tag}.xlsx]"
    def g(sheet,cell):
        k=f"{base}{sheet.upper()}'!{cell}"
        v=sol.get(k); 
        try: return v.value[0][0]
        except Exception: return v
    return g, time.time()-t0, sol
if __name__=='__main__':
    g,dt,sol=run({}, 'base')
    print('calc secs',round(dt,1))
    errs=[k for k,v in sol.items() if hasattr(v,'value') and any(isinstance(x,formulas.tokens.operand.XlError) or (isinstance(x,str) and x.startswith('#')) for x in v.value.ravel())]
    print('errors',len(errs), errs[:15])
    for r in range(19,29): print(g('Cover',f'A{r}'), g('Cover',f'B{r}'), g('Cover',f'C{r}'), g('Cover',f'D{r}'))
